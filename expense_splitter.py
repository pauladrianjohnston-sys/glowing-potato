#!/usr/bin/env python3
"""
HF Weekend Expense Splitter
─────────────────────────────────────────────────────────────────────────────
Workflow:
  1. Edit ATTENDEES and EXPENSES below (names, amounts, who paid).
  2. Run script → creates the Excel with a pre-filled Weights sheet.
  3. Open Excel → edit the Weights sheet (0.0 = excluded, 1.0 = full share,
     0.5 = half share, etc. — any positive number works proportionally).
  4. Re-run script → weights are read back, all other sheets refresh.
     Your weight edits survive every re-run.
"""

import os
from datetime import date
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.formatting.rule import CellIsRule

OUTPUT_FILE = os.path.join(os.path.dirname(__file__), "HF_Weekend_Expense_Splitter.xlsx")
CURRENCY = "A$"

# ── Attendees ──────────────────────────────────────────────────────────────────
ATTENDEES = [
    "Klaus", "Louise", "Alex", "Sam", "Jordan", "Taylor",
    "Morgan", "Jamie", "Casey", "Riley", "Drew", "Quinn",
    "Avery", "Blake", "Charlie", "Dylan",
]

# ── Expenses ───────────────────────────────────────────────────────────────────
# "splits" sets the INITIAL weights written to the Weights sheet on first run.
# After that, edit the Weights sheet directly in Excel — this field is ignored.
#   "even"                  → everyone starts at 1.0
#   {"Name": w, "default": 1.0}  → named people get w, rest get default
EXPENSES = [
    {
        "date": "Fri",
        "description": "Accommodation",
        "category": "Accommodation",
        "paid_by": "Alex",
        "amount": 960.00,
        "notes": "Full weekend, all 16",
        "splits": "even",
    },
    {
        "date": "Sat",
        "description": "Saturday Lunch",
        "category": "Food",
        "paid_by": "Sam",
        "amount": 240.00,
        "notes": "",
        "splits": "even",
    },
    {
        "date": "Sat",
        "description": "Saturday Drinks",
        "category": "Drinks",
        "paid_by": "Taylor",
        "amount": 320.00,
        "notes": "Klaus smaller share (~1/3)",
        "splits": {"Klaus": 0.3, "default": 1.0},
    },
    {
        "date": "Sat",
        "description": "Saturday Dinner",
        "category": "Food",
        "paid_by": "Morgan",
        "amount": 480.00,
        "notes": "",
        "splits": "even",
    },
    {
        "date": "Sun",
        "description": "Sunday Breakfast",
        "category": "Food",
        "paid_by": "Jamie",
        "amount": 280.00,
        "notes": "Klaus & Louise not present",
        "splits": {"Klaus": 0.0, "Louise": 0.0, "default": 1.0},
    },
    {
        "date": "Sun",
        "description": "Sunday Lunch",
        "category": "Food",
        "paid_by": "Casey",
        "amount": 360.00,
        "notes": "",
        "splits": "even",
    },
]

# ── Style helpers ──────────────────────────────────────────────────────────────
def fill(hex_color):
    return PatternFill("solid", fgColor=hex_color)

def fnt(bold=False, color="000000", size=11, italic=False):
    return Font(bold=bold, color=color, size=size, italic=italic, name="Calibri")

THIN = Side(style="thin")
BDR = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

def hdr(ws, row, col, value, bg="1F3864", fg="FFFFFF", bold=True, size=11, wrap=False):
    c = ws.cell(row=row, column=col, value=value)
    c.fill = fill(bg); c.font = fnt(bold=bold, color=fg, size=size)
    c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=wrap)
    c.border = BDR
    return c

# ── Weight helpers ─────────────────────────────────────────────────────────────
WEIGHTS_DATA_ROW = 5   # first person row in Weights sheet

def default_weights():
    """Convert EXPENSES 'splits' config → weight matrix for initial Excel."""
    w = {}
    for idx, exp in enumerate(EXPENSES):
        sp = exp.get("splits", "even")
        if sp == "even":
            w[idx] = {p: 1.0 for p in ATTENDEES}
        else:
            dw = sp.get("default", 1.0)
            w[idx] = {p: float(sp.get(p, dw)) for p in ATTENDEES}
    return w

def read_weights():
    """Read weights from existing Excel Weights sheet. Returns None if absent."""
    if not os.path.exists(OUTPUT_FILE):
        return None
    try:
        wb = openpyxl.load_workbook(OUTPUT_FILE, data_only=True)
        if "Weights" not in wb.sheetnames:
            return None
        ws = wb["Weights"]
        # Build person→row map
        person_row = {}
        for r in range(WEIGHTS_DATA_ROW, WEIGHTS_DATA_ROW + len(ATTENDEES) + 5):
            v = ws.cell(row=r, column=1).value
            if v and str(v).strip() in ATTENDEES:
                person_row[str(v).strip()] = r
        w = {}
        for idx in range(len(EXPENSES)):
            col = idx + 2
            w[idx] = {}
            for p in ATTENDEES:
                if p in person_row:
                    raw = ws.cell(row=person_row[p], column=col).value
                    try:
                        w[idx][p] = float(raw) if raw is not None else 0.0
                    except (ValueError, TypeError):
                        w[idx][p] = 0.0
                else:
                    w[idx][p] = 0.0
        print("  Weights read from existing Excel.")
        return w
    except Exception as e:
        print(f"  Warning: could not read weights ({e}), using defaults.")
        return None

# ── Core calculations ──────────────────────────────────────────────────────────
def compute_splits(weights):
    splits_matrix = {}
    paid_totals = {p: 0.0 for p in ATTENDEES}
    owed_totals = {p: 0.0 for p in ATTENDEES}
    for idx, exp in enumerate(EXPENSES):
        total = exp["amount"]
        paid_totals[exp["paid_by"]] = round(paid_totals[exp["paid_by"]] + total, 2)
        wts = weights[idx]
        total_weight = sum(wts.values())
        shares = {p: round(total * wts[p] / total_weight, 2) if total_weight else 0.0
                  for p in ATTENDEES}
        diff = round(total - sum(shares.values()), 2)
        if diff:
            for p in ATTENDEES:
                if shares[p] > 0:
                    shares[p] = round(shares[p] + diff, 2)
                    break
        splits_matrix[idx] = shares
        for p in ATTENDEES:
            owed_totals[p] = round(owed_totals[p] + shares[p], 2)
    net = {p: round(paid_totals[p] - owed_totals[p], 2) for p in ATTENDEES}
    return splits_matrix, paid_totals, owed_totals, net

def simplify_debts(net):
    cred = sorted([[v, k] for k, v in net.items() if v > 0.005], reverse=True)
    debt = sorted([[abs(v), k] for k, v in net.items() if v < -0.005], reverse=True)
    txns = []
    i = j = 0
    while i < len(cred) and j < len(debt):
        t = round(min(cred[i][0], debt[j][0]), 2)
        txns.append((debt[j][1], cred[i][1], t))
        cred[i][0] = round(cred[i][0] - t, 2)
        debt[j][0] = round(debt[j][0] - t, 2)
        if cred[i][0] < 0.01: i += 1
        if debt[j][0] < 0.01: j += 1
    return txns

# ── Sheet builders ─────────────────────────────────────────────────────────────
def build_overview(wb):
    ws = wb.create_sheet("Overview")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3
    ws.column_dimensions["B"].width = 72
    total = sum(e["amount"] for e in EXPENSES)

    for row, (txt, bg, fg, bold, h) in enumerate([
        (f"HF WEEKEND — EXPENSE SPLITTER  ({CURRENCY})", "1F3864", "FFFFFF", True, 44),
        (f"Generated {date.today().strftime('%d %b %Y')}  ·  {len(ATTENDEES)} attendees  ·  {len(EXPENSES)} expenses  ·  Total {CURRENCY}{total:,.2f}", "2F5496", "FFFFFF", False, 22),
    ], start=2):
        c = ws.cell(row=row, column=2, value=txt)
        c.fill = fill(bg); c.font = Font(bold=bold, color=fg, size=16 if bold else 11, name="Calibri")
        c.alignment = Alignment(horizontal="center", vertical="center")
        ws.row_dimensions[row].height = h

    lines = [
        ("HOW TO USE", "1F3864", "FFFFFF", True),
        ("1.  Edit ATTENDEES & EXPENSES in expense_splitter.py (names, amounts, who paid).", "EEF3FB", "000000", False),
        ("2.  Run script → Excel created with a pre-filled Weights sheet.", "EEF3FB", "000000", False),
        ("3.  Open Excel → go to the Weights sheet → enter each person's weight per expense.", "EEF3FB", "000000", False),
        ("4.  Re-run script → weights are read back, everything else refreshes.", "EEF3FB", "000000", False),
        ("", "FFFFFF", "000000", False),
        ("WEIGHTS SHEET RULES", "1F3864", "FFFFFF", True),
        (f"  0.0  = person not included in that expense (grey cell)", "FFFDE7", "000000", False),
        (f"  1.0  = standard full share", "FFFDE7", "000000", False),
        (f"  0.5  = half a standard share", "FFFDE7", "000000", False),
        (f"  Shares are proportional: e.g. ten 1.0s + two 0.5s → each 1.0 pays 1/11, each 0.5 pays 1/22", "FFFDE7", "000000", False),
        (f"  To split evenly among N people: give each of those N people any equal weight (e.g. all 1.0)", "FFFDE7", "000000", False),
    ]
    row = 5
    for txt, bg, fg, bold in lines:
        c = ws.cell(row=row, column=2, value=txt)
        c.fill = fill(bg); c.font = Font(bold=bold, color=fg, size=11, name="Calibri")
        c.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        ws.row_dimensions[row].height = 20
        row += 1

def build_attendees(wb):
    ws = wb.create_sheet("Attendees")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 46
    hdr(ws, 1, 1, "#", size=10); hdr(ws, 1, 2, "Name"); hdr(ws, 1, 3, "Notes")
    notes_map = {
        "Klaus": "Smaller share Sat Drinks; absent Sun Breakfast",
        "Louise": "Absent Sunday Breakfast",
    }
    for i, name in enumerate(ATTENDEES, 1):
        r = i + 1; bg = "F2F2F2" if i % 2 == 0 else "FFFFFF"
        for col in range(1, 4):
            ws.cell(row=r, column=col).fill = fill(bg)
            ws.cell(row=r, column=col).border = BDR
        ws.cell(row=r, column=1, value=i).alignment = Alignment(horizontal="center")
        ws.cell(row=r, column=2, value=name)
        nc = ws.cell(row=r, column=3, value=notes_map.get(name, ""))
        if name in notes_map:
            nc.font = Font(italic=True, color="444444", name="Calibri")

def build_expenses(wb):
    ws = wb.create_sheet("Expenses")
    ws.sheet_view.showGridLines = False
    cols = ["#", "Day", "Description", "Category", "Paid By", f"Amount ({CURRENCY})", "Notes"]
    widths = [5, 6, 28, 16, 16, 14, 42]
    for i, (c, w) in enumerate(zip(cols, widths), 1):
        ws.column_dimensions[get_column_letter(i)].width = w
        hdr(ws, 1, i, c)
    cat_colors = {"Accommodation": "BDD7EE", "Food": "E2EFDA",
                  "Drinks": "FCE4D6", "Transport": "FFF2CC", "Activities": "E8D5F5"}
    total = 0
    for idx, exp in enumerate(EXPENSES, 1):
        r = idx + 1; bg = cat_colors.get(exp["category"], "F2F2F2")
        for col, val in enumerate([idx, exp["date"], exp["description"], exp["category"],
                                    exp["paid_by"], exp["amount"], exp["notes"]], 1):
            c = ws.cell(row=r, column=col, value=val)
            c.fill = fill(bg); c.border = BDR
            if col == 6:
                c.number_format = f'"{CURRENCY}"#,##0.00'; c.alignment = Alignment(horizontal="right")
            if col == 7:
                c.font = Font(italic=True, color="555555", name="Calibri")
        total += exp["amount"]
    tr = len(EXPENSES) + 2
    for col in range(1, 8):
        ws.cell(row=tr, column=col).fill = fill("1F3864"); ws.cell(row=tr, column=col).border = BDR
    ws.cell(row=tr, column=5, value="TOTAL").font = fnt(bold=True, color="FFFFFF")
    c = ws.cell(row=tr, column=6, value=total)
    c.font = fnt(bold=True, color="FFFFFF"); c.number_format = f'"{CURRENCY}"#,##0.00'
    c.alignment = Alignment(horizontal="right")

def build_weights(wb, weights):
    ws = wb.create_sheet("Weights")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 16
    n_exp = len(EXPENSES)
    for idx in range(n_exp):
        ws.column_dimensions[get_column_letter(idx + 2)].width = 13

    # Title / instruction rows
    last_col = get_column_letter(n_exp + 1)
    ws.merge_cells(f"A1:{last_col}1")
    c = ws["A1"]
    c.value = f"SPLIT WEIGHTS — edit this sheet, then re-run expense_splitter.py to refresh"
    c.fill = fill("1F3864"); c.font = fnt(bold=True, color="FFFFFF", size=12)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 28

    ws.merge_cells(f"A2:{last_col}2")
    c = ws["A2"]
    c.value = ("0.0 = excluded  |  1.0 = full share  |  0.5 = half share  |  "
               "Any positive number works — shares are proportional to weights in each column")
    c.fill = fill("FFFDE7"); c.font = fnt(size=10, italic=True)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[2].height = 20

    # Row 3: expense name headers
    hdr(ws, 3, 1, "Person", "2F5496", wrap=True)
    for idx, exp in enumerate(EXPENSES):
        hdr(ws, 3, idx + 2, f"{exp['date']}: {exp['description']}", "2F5496", size=9, wrap=True)
    ws.row_dimensions[3].height = 38

    # Row 4: paid-by + amount sub-headers
    for col, label in [(1, "Paid by →")]:
        c = ws.cell(row=4, column=col, value=label)
        c.fill = fill("D9E1F2"); c.font = fnt(bold=True, size=9); c.border = BDR
    for idx, exp in enumerate(EXPENSES):
        c = ws.cell(row=4, column=idx + 2,
                    value=f"{exp['paid_by']}  {CURRENCY}{exp['amount']:,.0f}")
        c.fill = fill("D9E1F2"); c.font = fnt(bold=True, size=9)
        c.alignment = Alignment(horizontal="center"); c.border = BDR
    ws.row_dimensions[4].height = 18

    # Person weight rows (WEIGHTS_DATA_ROW = 5)
    data_range = f"B{WEIGHTS_DATA_ROW}:{get_column_letter(n_exp + 1)}{WEIGHTS_DATA_ROW + len(ATTENDEES) - 1}"
    for pidx, person in enumerate(ATTENDEES):
        r = WEIGHTS_DATA_ROW + pidx
        bg = "F2F2F2" if pidx % 2 == 0 else "FFFFFF"
        c = ws.cell(row=r, column=1, value=person)
        c.fill = fill("D9E1F2"); c.font = fnt(bold=True); c.border = BDR
        for idx in range(n_exp):
            w = weights[idx].get(person, 0.0)
            c = ws.cell(row=r, column=idx + 2, value=w)
            c.number_format = "0.0#"
            c.alignment = Alignment(horizontal="center"); c.border = BDR
            c.fill = fill("EEEEEE") if w == 0.0 else fill(bg)

    # Conditional formatting: green when > 0, grey when = 0
    ws.conditional_formatting.add(
        data_range,
        CellIsRule(operator="greaterThan", formula=["0"],
                   fill=PatternFill("solid", fgColor="C6EFCE")))
    ws.conditional_formatting.add(
        data_range,
        CellIsRule(operator="equal", formula=["0"],
                   fill=PatternFill("solid", fgColor="DDDDDD")))

def build_splits(wb, splits_matrix):
    ws = wb.create_sheet("Splits")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 16
    n_exp = len(EXPENSES)
    total_col = n_exp + 2
    for idx in range(n_exp):
        ws.column_dimensions[get_column_letter(idx + 2)].width = 13
    ws.column_dimensions[get_column_letter(total_col)].width = 14

    hdr(ws, 1, 1, "Person", "1F3864", wrap=True)
    for idx, exp in enumerate(EXPENSES):
        hdr(ws, 1, idx + 2, f"{exp['date']}: {exp['description']}", "2F5496", size=9, wrap=True)
    hdr(ws, 1, total_col, "Total Owed", "ED7D31", size=10)
    ws.row_dimensions[1].height = 42

    for row_lbl, getter in [(2, lambda e: e["paid_by"]), (3, lambda e: f"{CURRENCY}{e['amount']:,.0f}")]:
        ws.cell(row=row_lbl, column=1, value="Paid by:" if row_lbl == 2 else "Amount:").fill = fill("D9E1F2")
        ws.cell(row=row_lbl, column=1).font = fnt(bold=True, size=9)
        ws.cell(row=row_lbl, column=1).border = BDR
        for idx, exp in enumerate(EXPENSES):
            c = ws.cell(row=row_lbl, column=idx + 2, value=getter(exp))
            c.fill = fill("D9E1F2"); c.font = fnt(bold=True, size=9)
            c.alignment = Alignment(horizontal="center"); c.border = BDR
        ws.cell(row=row_lbl, column=total_col).fill = fill("FCE4D6")
        ws.cell(row=row_lbl, column=total_col).border = BDR

    for pidx, person in enumerate(ATTENDEES):
        r = pidx + 4; bg = "F2F2F2" if pidx % 2 == 0 else "FFFFFF"
        c = ws.cell(row=r, column=1, value=person)
        c.fill = fill("D9E1F2"); c.font = fnt(bold=True); c.border = BDR
        person_total = 0.0
        for eidx in range(n_exp):
            share = splits_matrix[eidx][person]
            c = ws.cell(row=r, column=eidx + 2)
            c.border = BDR
            if share == 0:
                c.value = "—"; c.fill = fill("EEEEEE")
                c.font = fnt(color="AAAAAA", italic=True)
                c.alignment = Alignment(horizontal="center")
            else:
                c.value = share; c.number_format = f'"{CURRENCY}"#,##0.00'
                c.alignment = Alignment(horizontal="right"); c.fill = fill(bg)
            person_total += share
        c = ws.cell(row=r, column=total_col, value=round(person_total, 2))
        c.number_format = f'"{CURRENCY}"#,##0.00'; c.alignment = Alignment(horizontal="right")
        c.fill = fill("FCE4D6"); c.font = fnt(bold=True); c.border = BDR

def build_balances(wb, paid_totals, owed_totals, net):
    ws = wb.create_sheet("Balances")
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDE", [22, 16, 16, 16, 30]):
        ws.column_dimensions[col].width = w
    for i, h in enumerate([f"Person", f"Total Paid ({CURRENCY})", f"Total Share ({CURRENCY})",
                            f"Net Balance ({CURRENCY})", "Status"], 1):
        hdr(ws, 1, i, h)
    for ridx, person in enumerate(sorted(ATTENDEES, key=lambda p: net[p], reverse=True), 1):
        r = ridx + 1; n = net[person]; bg_row = "F2F2F2" if ridx % 2 == 0 else "FFFFFF"
        if n > 0.01:   net_bg, net_fg, s_bg, status = "C6EFCE", "276221", "C6EFCE", f"Gets back {CURRENCY}{n:.2f}"
        elif n < -0.01: net_bg, net_fg, s_bg, status = "FFC7CE", "9C0006", "FFC7CE", f"Owes {CURRENCY}{abs(n):.2f}"
        else:           net_bg, net_fg, s_bg, status = "FFFF99", "7D6608", "FFFF99", "Settled"
        ws.cell(row=r, column=1, value=person).fill = fill("D9E1F2")
        ws.cell(row=r, column=1).font = fnt(bold=True); ws.cell(row=r, column=1).border = BDR
        for col, val in [(2, paid_totals[person]), (3, owed_totals[person])]:
            c = ws.cell(row=r, column=col, value=val)
            c.number_format = f'"{CURRENCY}"#,##0.00'; c.alignment = Alignment(horizontal="right")
            c.fill = fill(bg_row); c.border = BDR
        c = ws.cell(row=r, column=4, value=n)
        c.number_format = f'"{CURRENCY}"#,##0.00'; c.alignment = Alignment(horizontal="right")
        c.fill = fill(net_bg); c.font = fnt(bold=True, color=net_fg); c.border = BDR
        c = ws.cell(row=r, column=5, value=status)
        c.fill = fill(s_bg); c.font = fnt(bold=(n != 0)); c.border = BDR
    tr = len(ATTENDEES) + 2
    for col in range(1, 6):
        ws.cell(row=tr, column=col).fill = fill("1F3864"); ws.cell(row=tr, column=col).border = BDR
    ws.cell(row=tr, column=1, value="TOTAL").font = fnt(bold=True, color="FFFFFF")
    for col, val in [(2, sum(paid_totals.values())), (3, sum(owed_totals.values())),
                     (4, round(sum(net.values()), 2))]:
        c = ws.cell(row=tr, column=col, value=round(val, 2))
        c.number_format = f'"{CURRENCY}"#,##0.00'; c.alignment = Alignment(horizontal="right")
        c.font = fnt(bold=True, color="FFFFFF")

def build_settlement(wb, transactions):
    ws = wb.create_sheet("Settlement")
    ws.sheet_view.showGridLines = False
    for col, w in zip("ABCDEFG", [3, 5, 20, 8, 20, 14, 36]):
        ws.column_dimensions[col].width = w
    ws.merge_cells("B2:G2")
    c = ws["B2"]; c.value = "Recommended Settlement Plan — Minimum Transactions"
    c.fill = fill("1F3864"); c.font = Font(bold=True, color="FFFFFF", size=14, name="Calibri")
    c.alignment = Alignment(horizontal="center", vertical="center"); ws.row_dimensions[2].height = 34
    ws.merge_cells("B3:G3")
    c = ws["B3"]; c.value = f"{len(transactions)} payment(s) to fully settle all debts"
    c.fill = fill("2F5496"); c.font = Font(italic=True, color="FFFFFF", size=11, name="Calibri")
    c.alignment = Alignment(horizontal="center", vertical="center"); ws.row_dimensions[3].height = 22
    for i, h in enumerate(["#", "Pays", "→", "To", f"Amount ({CURRENCY})", "Bank ref"], 2):
        hdr(ws, 5, i, h, "2F5496")
    for tidx, (sender, receiver, amount) in enumerate(transactions, 1):
        r = tidx + 5; bg = "F2F2F2" if tidx % 2 == 0 else "FFFFFF"
        for col, val in [(2, tidx), (3, sender), (4, "→"), (5, receiver), (6, amount),
                         (7, f"HF Weekend {sender}")]:
            c = ws.cell(row=r, column=col, value=val)
            c.fill = fill(bg); c.border = BDR
            if col == 4:
                c.alignment = Alignment(horizontal="center")
                c.font = Font(bold=True, size=14, color="2F5496", name="Calibri")
            elif col == 6:
                c.number_format = f'"{CURRENCY}"#,##0.00'; c.alignment = Alignment(horizontal="right")
                c.font = fnt(bold=True)
            elif col in (3, 5): c.font = fnt(bold=True)
            elif col == 7: c.font = fnt(italic=True, color="666666")

def build_whatsapp(wb, transactions):
    ws = wb.create_sheet("WhatsApp")
    ws.sheet_view.showGridLines = False
    ws.column_dimensions["A"].width = 3; ws.column_dimensions["B"].width = 80
    total = sum(e["amount"] for e in EXPENSES)
    lines = [
        "─────────────────────────────────",
        "HF WEEKEND — EXPENSES & SETTLEMENT",
        "─────────────────────────────────",
        "",
        f"Total spend: {CURRENCY}{total:,.2f}  |  {len(ATTENDEES)} people",
        "",
        "Expenses:",
    ]
    for exp in EXPENSES:
        lines.append(f"  {exp['date']}  {exp['description']:<22} {CURRENCY}{exp['amount']:>8,.2f}  (paid by {exp['paid_by']})")
    lines += [
        "",
        "Split adjustments:",
        "  • Sat Drinks: Klaus had a small share (~1/3 of standard)",
        "  • Sun Breakfast: Klaus & Louise not present — split 14 ways",
        "",
        "Payments needed:",
    ]
    for sender, receiver, amount in transactions:
        lines.append(f"  {sender} pays {receiver}: {CURRENCY}{amount:,.2f}")
    lines += [
        "",
        "Please pay by bank transfer. Use your name as the reference.",
        "─────────────────────────────────",
    ]
    hdr(ws, 1, 2, "WhatsApp Message  —  copy the yellow block below and paste into your chat")
    for ridx, line in enumerate(lines, 2):
        c = ws.cell(row=ridx, column=2, value=line)
        c.font = Font(name="Courier New", size=10)
        c.alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[ridx].height = 16
    block = len(lines) + 4
    ws.merge_cells(f"B{block}:B{block + 40}")
    c = ws[f"B{block}"]; c.value = "\n".join(lines)
    c.font = Font(name="Courier New", size=10)
    c.alignment = Alignment(wrap_text=True, vertical="top")
    c.fill = fill("FFFDE7")
    for r in range(block, block + 41):
        ws.row_dimensions[r].height = 15

# ── Main ───────────────────────────────────────────────────────────────────────
def main():
    print("HF Weekend Expense Splitter")
    print("=" * 42)
    weights = read_weights() or default_weights()
    splits_matrix, paid_totals, owed_totals, net = compute_splits(weights)
    transactions = simplify_debts(net)
    total = sum(e["amount"] for e in EXPENSES)
    print(f"  Attendees : {len(ATTENDEES)}")
    print(f"  Expenses  : {len(EXPENSES)}  (total {CURRENCY}{total:,.2f})")
    print(f"  Settlement: {len(transactions)} transaction(s)\n")

    wb = openpyxl.Workbook(); wb.remove(wb.active)
    build_overview(wb)
    build_attendees(wb)
    build_expenses(wb)
    build_weights(wb, weights)
    build_splits(wb, splits_matrix)
    build_balances(wb, paid_totals, owed_totals, net)
    build_settlement(wb, transactions)
    build_whatsapp(wb, transactions)
    wb.save(OUTPUT_FILE)
    print(f"Saved: {OUTPUT_FILE}\n")

    print("Settlement plan:")
    for sender, receiver, amount in transactions:
        print(f"  {sender:<14} pays  {receiver:<14}  {CURRENCY}{amount:,.2f}")
    print("\nNet balances:")
    for p in sorted(ATTENDEES, key=lambda x: net[x], reverse=True):
        n = net[p]
        print(f"  {p:<14} {'gets back' if n > 0 else 'owes    '}  {CURRENCY}{abs(n):,.2f}")

if __name__ == "__main__":
    main()
